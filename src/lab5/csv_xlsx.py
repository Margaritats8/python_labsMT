import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Converts CSV to XLSX.

     The first line of the CSV is the header. The sheet is named “Sheet1”. 
     The width of each column is equal to the length of the longest value, but not less than 8 characters.
     Values are written as strings (types are not guessed).

     Raises:
     FileNotFoundError: no input CSV.
     ValueError: invalid extension, empty CSV.
    """
    src = Path(csv_path)
    dst = Path(xlsx_path)
    if src.suffix.lower() != ".csv":
        raise ValueError(f"Invalid file type: expected .csv, received '{src.suffix}'")
    if dst.suffix.lower() != ".xlsx":
        raise ValueError(f"Invalid file type: expected .xlsx, received '{dst.suffix}'")
    if not src.is_file():
        raise FileNotFoundError(f"The file was not found: {csv_path}")

    rows = []
    with open(csv_path, encoding="utf-8", newline="") as f:
        for row in csv.reader(f):
            if row:
                rows.append(row)

    if len(rows) < 2:
        raise ValueError("CSV пустой или в нём нет данных (только заголовок)")

    widths = {}
    for row in rows:
        for index, cell in enumerate(row):
            if len(cell) > widths.get(index, 0): 
                widths[index] = len(cell)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)

    for index, width in widths.items():
        letter = get_column_letter(index + 1)
        ws.column_dimensions[letter].width = max(8, width) 

    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)

csv_to_xlsx("data/samples/people.csv", "data/out/people.xlsx")